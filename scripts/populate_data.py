import os
import sys
import json
import openpyxl
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

if sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass


def parse_teacher_docx(doc_path):
    doc = docx.Document(doc_path)
    table = doc.tables[0]
    teachers_raw = []
    
    # Header is rows 0 and 1
    for r in range(2, len(table.rows)):
        row_cells = [cell.text.strip() for cell in table.rows[r].cells]
        no = row_cells[0]
        nama = row_cells[1]
        mapel = row_cells[2]
        c_vii = row_cells[3]
        c_viii = row_cells[4]
        c_ix = row_cells[5]
        
        teachers_raw.append({
            "no": no,
            "nama": nama,
            "mapel": mapel,
            "vii": c_vii,
            "viii": c_viii,
            "ix": c_ix
        })
    return teachers_raw

def parse_student_excels(excel_files):
    all_students = []
    
    for f in excel_files:
        wb = openpyxl.load_workbook(f)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # Headers in row 1
            for r in range(2, ws.max_row + 1):
                no = ws.cell(r, 1).value
                nama = ws.cell(r, 2).value
                nisn = str(ws.cell(r, 3).value or '').strip()
                nik = str(ws.cell(r, 4).value or '').replace("'", '').strip()
                tempat_lahir = ws.cell(r, 5).value or ''
                tgl_lahir = str(ws.cell(r, 6).value or '').split(' ')[0]
                rombels = ws.cell(r, 7).value or sheet_name
                umur = ws.cell(r, 8).value or ''
                status = ws.cell(r, 9).value or 'Aktif'
                jk = ws.cell(r, 10).value or ''
                alamat = ws.cell(r, 11).value or ''
                telepon = ws.cell(r, 12).value or ''
                
                if nama:
                    # Clean up class name
                    clean_kelas = sheet_name
                    if "IX-A" in sheet_name or "9" in sheet_name:
                        clean_kelas = "IX-A"
                    elif "VIII-A" in sheet_name:
                        clean_kelas = "VIII-A"
                    elif "VIII B" in sheet_name or "VIII-B" in sheet_name:
                        clean_kelas = "VIII-B"
                    elif "VII-A" in sheet_name:
                        clean_kelas = "VII-A"
                    elif "VII-B" in sheet_name:
                        clean_kelas = "VII-B"
                        
                    all_students.append({
                        "no": str(no),
                        "nama": str(nama).strip(),
                        "nisn": nisn,
                        "nik": nik,
                        "tempat_lahir": str(tempat_lahir).strip(),
                        "tanggal_lahir": tgl_lahir,
                        "kelas": clean_kelas,
                        "rombel_full": str(rombels).strip(),
                        "umur": str(umur).strip(),
                        "status": str(status).strip(),
                        "jk": str(jk).strip(),
                        "alamat": str(alamat).strip(),
                        "telepon": str(telepon).strip()
                    })
    return all_students

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_word_document(teachers, students, output_filename):
    doc = docx.Document()
    
    # Page setup - Normal Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("DATA GURU DAN SISWA PERPUSTAKAAN EL-HUSNA")
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Tahun Ajaran 2025/2026 - Semester Ganjil")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    doc.add_paragraph() # Spacing

    # Section 1: Data Guru
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("I. DATA GURU DAN PENGAMPU MATA PELAJARAN")
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_g_desc = doc.add_paragraph("Berikut adalah daftar guru pengampu mata pelajaran untuk kelas VII, VIII, dan IX:")
    p_g_desc.runs[0].font.size = Pt(10)

    # Teacher Table
    table_g = doc.add_table(rows=1, cols=6)
    table_g.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Headers
    hdr_cells = table_g.rows[0].cells
    hdr_titles = ['No', 'Nama Pemeriksa / Guru', 'Mata Pelajaran', 'Kelas VII', 'Kelas VIII', 'Kelas IX']

    for idx, title in enumerate(hdr_titles):
        hdr_cells[idx].text = title
        set_cell_background(hdr_cells[idx], "1E3A8A")
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(9.5)

    for row_idx, g in enumerate(teachers):
        row_cells = table_g.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        
        vals = [g["no"], g["nama"], g["mapel"], g["vii"], g["viii"], g["ix"]]
        aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
        
        for idx, val in enumerate(vals):
            row_cells[idx].text = val
            set_cell_background(row_cells[idx], bg_color)
            set_cell_margins(row_cells[idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[idx].paragraphs[0]
            p.alignment = aligns[idx]
            for r in p.runs:
                r.font.size = Pt(9)

    doc.add_page_break()

    # Section 2: Data Siswa
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("II. DATA SISWA TERDAFTAR")
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Class summary
    classes = ["VII-A", "VII-B", "VIII-A", "VIII-B", "IX-A"]
    class_students = {c: [s for s in students if s["kelas"] == c] for c in classes}

    doc.add_paragraph("Rekapitulasi jumlah siswa terdaftar berdasarkan rombel / kelas:")
    
    # Summary Table
    t_sum = doc.add_table(rows=1, cols=3)
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_hdr = t_sum.rows[0].cells
    s_hdr[0].text = "Tingkat / Kelas"
    s_hdr[1].text = "Jumlah Siswa"
    s_hdr[2].text = "Status Kategori"
    
    for c in s_hdr:
        set_cell_background(c, "2563EB")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(9.5)

    total_all = 0
    for idx, c_name in enumerate(classes):
        cnt = len(class_students[c_name])
        total_all += cnt
        row = t_sum.add_row().cells
        bg = "F1F5F9" if idx % 2 == 1 else "FFFFFF"
        row[0].text = f"Kelas {c_name}"
        row[1].text = f"{cnt} Siswa"
        row[2].text = "Aktif"
        
        for i in range(3):
            set_cell_background(row[i], bg)
            set_cell_margins(row[i], top=80, bottom=80, left=100, right=100)
            p = row[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)

    # Total row
    row_tot = t_sum.add_row().cells
    row_tot[0].text = "TOTAL KESELURUHAN"
    row_tot[1].text = f"{total_all} Siswa"
    row_tot[2].text = "Aktif"
    for i in range(3):
        set_cell_background(row_tot[i], "DBEAFE")
        set_cell_margins(row_tot[i], top=100, bottom=100, left=100, right=100)
        p = row_tot[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.5)

    doc.add_paragraph() # spacing

    # Student tables per class
    for c_name in classes:
        h_c = doc.add_heading(level=2)
        r_c = h_c.add_run(f"Daftar Siswa Kelas {c_name} ({len(class_students[c_name])} Siswa)")
        r_c.font.size = Pt(11)
        r_c.font.bold = True
        r_c.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        t_s = doc.add_table(rows=1, cols=7)
        t_s.alignment = WD_TABLE_ALIGNMENT.CENTER
        s_hdr_cells = t_s.rows[0].cells
        headers = ['No', 'Nama Lengkap', 'NISN', 'NIK', 'L/P', 'Tempat, Tgl Lahir', 'Status']

        for i, h_text in enumerate(headers):
            s_hdr_cells[i].text = h_text
            set_cell_background(s_hdr_cells[i], "1E3A8A")
            set_cell_margins(s_hdr_cells[i], top=100, bottom=100, left=80, right=80)
            p = s_hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)

        for s_idx, st in enumerate(class_students[c_name]):
            r_cells = t_s.add_row().cells
            bg = "F8FAFC" if s_idx % 2 == 1 else "FFFFFF"
            
            ttl = f"{st['tempat_lahir']}, {st['tanggal_lahir']}" if st['tempat_lahir'] else st['tanggal_lahir']
            jk_short = "L" if "Laki" in st['jk'] else ("P" if "Perempuan" in st['jk'] else st['jk'])
            
            vals = [str(s_idx + 1), st['nama'], st['nisn'], st['nik'], jk_short, ttl, st['status']]
            aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER]

            for i in range(7):
                r_cells[i].text = vals[i]
                set_cell_background(r_cells[i], bg)
                set_cell_margins(r_cells[i], top=60, bottom=60, left=80, right=80)
                p = r_cells[i].paragraphs[0]
                p.alignment = aligns[i]
                for r in p.runs:
                    r.font.size = Pt(8.5)
        
        doc.add_paragraph() # Spacing

    doc.save(output_filename)
    print(f"✅ Word document generated successfully: {output_filename}")

def main():
    excel_files = [
        'Daftar_Siswa_2025_2026_Ganjil_1766800079845.xlsx',
        'Daftar_Siswa_2025_2026_Ganjil_1766800095752.xlsx',
        'Daftar_Siswa_2025_2026_Ganjil_1766800109838.xlsx'
    ]
    doc_path = 'data guru.docx'

    teachers = parse_teacher_docx(doc_path)
    students = parse_student_excels(excel_files)

    print(f"Extracted {len(teachers)} teacher entries and {len(students)} student entries.")

    # Save to JSON for Prisma Seed script
    export_data = {
        "teachers": teachers,
        "students": students
    }

    with open('db/students_teachers_data.json', 'w', encoding='utf-8') as f:
        json.dump(export_data, f, ensure_ascii=False, indent=2)
    print("✅ Exported json to db/students_teachers_data.json")

    # Generate Word Document
    create_word_document(teachers, students, 'data guru.docx')
    create_word_document(teachers, students, 'Data_Siswa_dan_Guru.docx')

if __name__ == '__main__':
    main()
